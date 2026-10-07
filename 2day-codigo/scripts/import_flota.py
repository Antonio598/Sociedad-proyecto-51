"""Importa el catálogo de flotilla desde FLOTILLA 2DAY: enriquece `unidades`
(serie, motor, placa, año, marca, descripción, operador asignado) y puebla
`remolques` con las banderas derivadas (es_dolly, usa_combustible).

Uso (desde backend/, venv activado):
    python -m scripts.import_flota "C:\\ruta\\FLOTILLA 2DAY 2026 MEX.xlsx"
"""

import sys
import warnings

import openpyxl
from sqlalchemy import select

from app.db import SessionLocal
from app.models import Remolque, TipoUnidad, Unidad

warnings.filterwarnings("ignore")


def txt(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def clave(v):
    if v is None:
        return None
    s = "".join(str(v).strip().upper().split())
    if len(s) < 2 or s.startswith("#") or s == "TOTAL":
        return None
    return s


def ent(v):
    try:
        return int(float(v)) if v is not None and str(v).strip() != "" else None
    except (ValueError, TypeError):
        return None


def tipo_de(clv: str) -> TipoUnidad:
    return TipoUnidad.TRACTO if clv.startswith("T") else TipoUnidad.CAMION


def cell(r, idx):
    return r[idx] if idx < len(r) else None


def importar_tractos(ws, ses) -> int:
    n = 0
    for r in ws.iter_rows(min_row=6, max_row=56, values_only=True):
        clv = clave(cell(r, 0))
        if not clv:
            continue
        u = ses.execute(select(Unidad).where(Unidad.clave == clv)).scalar_one_or_none()
        if u is None:
            u = Unidad(clave=clv)
            ses.add(u)
        u.tipo = tipo_de(clv)
        u.serie = txt(cell(r, 1))
        u.operador_asignado = txt(cell(r, 2))
        u.motor = txt(cell(r, 3))
        u.placa = txt(cell(r, 4))
        u.placas_nuevas = txt(cell(r, 5))
        u.anio = ent(cell(r, 6))
        u.marca = txt(cell(r, 7))
        u.descripcion = txt(cell(r, 8))
        u.usa_remolque = (u.tipo == TipoUnidad.TRACTO)
        u.activo = True
        n += 1
    ses.flush()
    return n


def importar_remolques(ws, ses) -> tuple[int, int]:
    """UPSERT por `eco`. Devuelve (nuevos, actualizados).

    ANTES hacía `ses.query(Remolque).delete()` y volvía a insertarlo todo. Eso era una
    bomba: hoy la operación apunta a los remolques por id y por texto, así que borrarlos
    en cada corrida (a) chocaría con la FK de SolicitudRecarga.remolque_id o dejaría
    huérfano el consumo de termo ya capturado, (b) cambiaría los ids a los que apunta
    AsignacionViaje.remolque_ids, y (c) borraría el titular que se asigna desde el panel
    y que NO viene en este Excel. Se cambió a upsert, el mismo patrón que ya usaba
    importar_tractos.

    Tampoco se borra lo que el Excel no liste: un remolque puede existir y estar
    consumiendo aunque falte en el archivo (se verificó: 531811 y 531818 cargan diésel
    y no aparecen). Darlo de baja en silencio es justo lo que no se debe hacer.
    """
    nuevos = actualizados = 0
    # Sin max_row: antes leía 6..65 (60 filas) y el catálogo real es mayor, así que
    # silenciosamente se dejaba fuera a los últimos remolques del archivo.
    for r in ws.iter_rows(min_row=6, values_only=True):
        eco = clave(cell(r, 0))
        if not eco:
            continue
        desc = txt(cell(r, 8))
        s_thermo = txt(cell(r, 3))
        if s_thermo and s_thermo.upper() in ("S/N", "SN", "N/A"):
            s_thermo = None
        es_dolly = eco.startswith("D-") or eco.startswith("D0") or (desc is not None and "DOLLY" in desc.upper())
        refrigerado = desc is not None and ("REFRIGERAD" in desc.upper() or "C/EQREF" in desc.upper() or "EQREF" in desc.upper())
        usa_comb = (not es_dolly) and (s_thermo is not None or refrigerado)

        x = ses.execute(select(Remolque).where(Remolque.eco == eco)).scalar_one_or_none()
        if x is None:
            x = Remolque(eco=eco)
            ses.add(x)
            nuevos += 1
        else:
            actualizados += 1
        x.eco_nuevo = txt(cell(r, 1))
        x.serie = txt(cell(r, 2))
        x.serie_thermo = s_thermo
        x.placa = txt(cell(r, 4))
        x.placas_nuevas = txt(cell(r, 5))
        x.anio = ent(cell(r, 6))
        x.marca = txt(cell(r, 7))
        x.descripcion = desc
        x.es_dolly = es_dolly
        x.usa_combustible = usa_comb
        # NO se tocan: `operador_asignado_id` (lo fija el coordinador desde el panel) ni
        # `activo` (para no resucitar un remolque dado de baja a propósito).
    ses.flush()
    return nuevos, actualizados


def main(path: str, dry: bool = False) -> None:
    from sqlalchemy import func

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ses = SessionLocal()
    antes_rem = ses.scalar(select(func.count()).select_from(Remolque)) or 0
    hojas = {h.upper(): h for h in wb.sheetnames}
    ntract = importar_tractos(wb[hojas["TRACTOS"]], ses)
    nuevos, actualizados = importar_remolques(wb[hojas["REMOLQUES"]], ses)

    if dry:
        ses.rollback()
        print("=== SIMULACIÓN (no se escribió nada) ===")
        print(f"  unidades que se tocarían : {ntract}")
        print(f"  remolques que se crearían: {nuevos}")
        print(f"  remolques que se actualizarían: {actualizados}")
        print(f"  remolques en la BD que el Excel NO lista: {max(0, antes_rem - actualizados)}"
              "  (se CONSERVAN: un remolque puede estar consumiendo aunque falte en el archivo)")
        ses.close()
        return

    ses.commit()
    total_uni = ses.scalar(select(func.count()).select_from(Unidad))
    total_rem = ses.scalar(select(func.count()).select_from(Remolque))
    dolly = ses.scalar(select(func.count()).select_from(Remolque).where(Remolque.es_dolly.is_(True)))
    con_comb = ses.scalar(select(func.count()).select_from(Remolque).where(Remolque.usa_combustible.is_(True)))
    print("=== CATÁLOGO DE FLOTILLA IMPORTADO ===")
    print(f"  unidades (enriquecidas/creadas): {ntract}  | total en BD: {total_uni}")
    print(f"  remolques: {nuevos} nuevos, {actualizados} actualizados  | total en BD: {total_rem}")
    print(f"    (dolly sin combustible: {dolly}, con termo/combustible: {con_comb})")
    no_listados = max(0, total_rem - nuevos - actualizados)
    if no_listados:
        print(f"  {no_listados} remolques de la BD no vienen en el Excel y se CONSERVARON.")
    ses.close()


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if not args:
        print('Uso: python -m scripts.import_flota "ruta\\FLOTILLA 2DAY 2026 MEX.xlsx" [--dry-run]')
        sys.exit(1)
    main(args[0], dry="--dry-run" in sys.argv)
