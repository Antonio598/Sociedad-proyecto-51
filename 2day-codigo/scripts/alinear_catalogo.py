"""Deja el catálogo alineado con los dos maestros del cliente: HOLOGRAMA y PLACAS.

Decisión del dueño (16-sep-2026): "se borrarán todas las unidades excepto las que
pertenecen a los archivos".

QUÉ HACE Y QUÉ NO. **No borra nada, nunca.** Y no es una licencia que se tome el script:
las 22 claves foráneas que apuntan a `unidades` y `remolques` son todas `NO ACTION`, así
que Postgres RECHAZA el DELETE de cualquier activo con un viaje, una carga o un asiento
detrás. "Borrar" sólo puede significar `activo = False`, que además es reversible con un
UPDATE que este script imprime.

UN REMOLQUE TIENE DOS NOMBRES. La flota se renumeró y el cambio está registrado en
`Remolque.eco_nuevo` (28 remolques: `5311802` pasó a ser `531812`, `53139` pasó a ser
`531810`). La renumeración es arbitraria y NO se puede deducir del texto: hay que leerla
de la columna. De ahí la regla que gobierna todo este script:

    un activo sobrevive si CUALQUIERA de sus nombres aparece en algún maestro.

Decidirlo por nombre en vez de por activo daría de baja a un remolque que el maestro sí
lista, sólo porque lo lista por su nombre nuevo. Es el error que este script NO comete.

Tampoco se adivinan parejas por parecido. Un `531806` en el maestro que no esté en la base
se informa para que lo mire una persona, no se empareja solo: el parecido de los textos
apuntaba a `5311806`, y la columna dice que `5311806` es en realidad `531816`. El
importador propone; no inventa.

LA REGLA QUE PROTEGE, heredada de `depurar_contra_hologramas.py` y ampliada: un activo con
consumo NO se da de baja aunque el maestro no lo liste, porque dejaría litros sin dueño. Ya
pasó con 5311801 y 5311808. Aquí se mira el libro mayor (`AsientoConsumo`) **y además** las
cargas del proveedor, que el script viejo no consultaba.

Las altas pasan por `catalogo.aplicar_propuesta`, que es el ÚNICO camino de alta que
además siembra el alias; los otros seis sitios que crean activos no lo hacen y dejan al
activo invisible para el resolvedor del proveedor.

Uso:
    python -m scripts.alinear_catalogo                 simula, no escribe nada
    python -m scripts.alinear_catalogo --aplicar       aplica altas y bajas
    python -m scripts.alinear_catalogo --aplicar RUTA_HOLOGRAMA RUTA_PLACAS
"""

import io as _io
import os
import sys
from datetime import datetime, timezone

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from sqlalchemy import select

from app import catalogo
from app.catalogo import norm_eco, norm_etiqueta
from app.db import SessionLocal
from app.models import AsientoConsumo, CargaProveedor, Remolque, Unidad, Viaje

DL = os.path.expanduser("~/Downloads")
DEF_HOLO = os.path.join(DL, "HOLOGRAMA.xlsx")
DEF_PLAC = os.path.join(DL, "PLACAS.xlsx")

CABECERAS = {"UNIDAD", "REMOLQUE", "PLACAS", "HOLOGRAMA", "ECO"}


class Propuesta:
    """Lo mínimo que `catalogo.aplicar_propuesta` necesita leer y escribir.

    Es duck-typing deliberado: se reutiliza la función de verdad —la que el endpoint del
    panel usa— en vez de repetir aquí la lógica de alta y arriesgarse a que las dos se
    separen con el tiempo.
    """

    def __init__(self, accion, entidad, eco, placa=None, entidad_id=None):
        self.estado = "pendiente"
        self.accion = accion
        self.entidad = entidad          # 'unidad' | 'remolque'
        self.entidad_id = entidad_id
        self.eco_texto = eco
        self.placa_texto = placa
        self.aplicada_por_id = None
        self.aplicada_en = None


# ── lectura de los maestros ────────────────────────────────────────────────────

def _hoja(path):
    with open(path, "rb") as fh:
        wb = openpyxl.load_workbook(_io.BytesIO(fh.read()), data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    ws.reset_dimensions()      # el <dimension> declarado recortaba la lectura en silencio
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _celda(fila, i):
    if i >= len(fila) or fila[i] is None:
        return ""
    v = str(fila[i]).strip()
    return "" if v.upper() in CABECERAS else v


def leer_holograma(path):
    """[(eco_crudo, codigo_crudo)]. El maestro mezcla unidades y remolques en una columna."""
    return [(e, _celda(f, 2)) for f in _hoja(path) if (e := _celda(f, 1))]


def leer_placas(path):
    """(unidades, remolques), cada una [(eco_crudo, placa_cruda)]."""
    uni, rem = [], []
    for fila in _hoja(path):
        if e := _celda(fila, 2):
            uni.append((e, _celda(fila, 3)))
        if e := _celda(fila, 5):
            rem.append((e, _celda(fila, 6)))
    return uni, rem


def es_unidad(eco):
    """T### y C### son tracto y camión. Lo demás —numérico o D-0x— es remolque."""
    t = norm_eco(eco)
    return bool(t) and t[0] in ("T", "C") and not t.isdigit()


def nombres(obj, es_uni):
    """Todos los textos por los que se conoce al activo. Un remolque tiene dos."""
    if es_uni:
        return {norm_eco(obj.clave)} - {""}
    return {norm_eco(obj.eco), norm_eco(obj.eco_nuevo)} - {""}


# ── la regla que protege ───────────────────────────────────────────────────────

def consumo(db, obj, es_uni):
    """(nº de evidencias, litros). Mira el libro mayor Y las cargas del proveedor.

    Son dos fuentes distintas a propósito: el libro mayor es lo ya contabilizado, y las
    cargas son lo que el proveedor facturó aunque todavía no se haya asentado. Un activo
    con cualquiera de las dos tiene litros a su nombre y no se toca.
    """
    col_a = AsientoConsumo.unidad_id if es_uni else AsientoConsumo.remolque_id
    col_c = CargaProveedor.unidad_id if es_uni else CargaProveedor.remolque_id
    asientos = db.query(AsientoConsumo).filter(col_a == obj.id).all()
    cargas = db.query(CargaProveedor).filter(col_c == obj.id,
                                             CargaProveedor.vigente.is_(True)).all()
    return (len(asientos) + len(cargas),
            sum(float(a.litros or 0) for a in asientos)
            + sum(float(c.litros or 0) for c in cargas))


def titulo(t):
    print(f"\n{t}\n{'─' * len(t)}")


def main(aplicar, ruta_holo, ruta_plac):
    holo = leer_holograma(ruta_holo)
    plac_u, plac_r = leer_placas(ruta_plac)

    m_uni, m_rem = {}, {}
    for e, _ in holo:                       # el maestro de hologramas no trae placa
        (m_uni if es_unidad(e) else m_rem).setdefault(norm_eco(e), "")
    for e, p in plac_u:
        m_uni[norm_eco(e)] = p or m_uni.get(norm_eco(e), "")
    for e, p in plac_r:
        m_rem[norm_eco(e)] = p or m_rem.get(norm_eco(e), "")
    m_uni.pop("", None)
    m_rem.pop("", None)

    print("MAESTROS")
    print(f"  {os.path.basename(ruta_holo)}: {len(holo)} filas")
    print(f"  {os.path.basename(ruta_plac)}: {len(plac_u)} unidades, {len(plac_r)} remolques")
    print(f"  unión -> {len(m_uni)} unidades, {len(m_rem)} remolques")

    db = SessionLocal()
    unidades = db.execute(select(Unidad)).scalars().all()
    remolques = db.execute(select(Remolque)).scalars().all()
    activos = [(u, True) for u in unidades] + [(r, False) for r in remolques]

    # Un solo índice de todos los nombres conocidos, para que una alta no duplique a un
    # activo que ya existe con su otro nombre.
    conocidos = set()
    for obj, es_uni in activos:
        conocidos |= nombres(obj, es_uni)

    ren = [r for r in remolques if norm_eco(r.eco_nuevo) and norm_eco(r.eco_nuevo) != norm_eco(r.eco)]
    print(f"  remolques con dos nombres (renumerados): {len(ren)}")

    # ── 1. ALTAS ───────────────────────────────────────────────────────────────
    altas = ([("unidad", e, m_uni[e]) for e in sorted(m_uni) if e not in conocidos]
             + [("remolque", e, m_rem[e]) for e in sorted(m_rem) if e not in conocidos])
    titulo(f"1. ALTAS ({len(altas)}) · el maestro los lista y no existen con ningún nombre")
    for ent, eco, placa in altas:
        print(f"   {ent:9} {eco:<12} placa {placa or '(sin placa en el maestro)'}")

    # ── 2. BAJAS ───────────────────────────────────────────────────────────────
    # Por ACTIVO, no por nombre: sobrevive si cualquiera de sus nombres está listado.
    bajas, protegidos = [], []
    for obj, es_uni in activos:
        if not obj.activo:
            continue
        maestro = m_uni if es_uni else m_rem
        if nombres(obj, es_uni) & set(maestro):
            continue
        n, lts = consumo(db, obj, es_uni)
        nv = (db.query(Viaje).filter(Viaje.unidad_id == obj.id).count() if es_uni
              else db.query(Viaje).filter(Viaje.remolque_thermo == obj.eco).count())
        ent = "unidad" if es_uni else "remolque"
        eco = obj.clave if es_uni else obj.eco
        (protegidos if n else bajas).append((ent, obj, eco, nv, n, lts))

    titulo(f"2. BAJAS ({len(bajas)}) · ningún maestro los lista y no tienen consumo")
    print("   `activo = False`. Conservan viajes, asientos y etiquetas. Reversible.")
    for ent, obj, eco, nv, _, _ in sorted(bajas, key=lambda x: -x[3]):
        extra = "  DOLLY" if ent == "remolque" and obj.es_dolly else ""
        otro = f"  (también {obj.eco_nuevo})" if ent == "remolque" and obj.eco_nuevo else ""
        print(f"   {ent:9} {eco:<12} {nv:>5} viajes que se conservan{extra}{otro}")

    titulo(f"   NO SE TOCAN ({len(protegidos)}) · tienen litros a su nombre")
    for ent, obj, eco, nv, n, lts in sorted(protegidos, key=lambda x: -x[5]):
        print(f"   {ent:9} {eco:<12} {n:>4} evidencias · {lts:>10,.0f} L")

    # ── 2b. REACTIVACIONES ─────────────────────────────────────────────────────
    # La decision dice que los activos de los archivos SON la flota. Un activo dado de
    # baja en una depuracion anterior que hoy vuelve a aparecer en un maestro tiene que
    # volver, o quedaria fuera de las asignaciones sin que nadie lo note.
    revivir = []
    for obj, es_uni in activos:
        if obj.activo:
            continue
        maestro = m_uni if es_uni else m_rem
        if nombres(obj, es_uni) & set(maestro):
            revivir.append(("unidad" if es_uni else "remolque", obj,
                            obj.clave if es_uni else obj.eco))
    titulo(f"2b. REACTIVACIONES ({len(revivir)}) · estaban de baja y el maestro los lista")
    for ent, obj, eco in revivir:
        otro = f"  (tambien {obj.eco_nuevo})" if ent == "remolque" and obj.eco_nuevo else ""
        print(f"   {ent:9} {eco:<12}{otro}")

    # ── 3. PARA QUE LO MIRE UNA PERSONA ────────────────────────────────────────
    # Un código del maestro sin activo y un activo sin código pueden ser el mismo camión
    # renumerado. NO se empareja por parecido: la renumeración real es arbitraria y una
    # pareja mal puesta manda el diésel de un remolque a otro sin dejar rastro.
    sueltos_base = [e for _, _, e, _, _, _ in bajas + protegidos]
    if altas and sueltos_base:
        titulo("3. REVISAR A MANO · ¿alguno de estos es el mismo activo renumerado?")
        print("   Si lo es, NO se da de alta: se le pone el nombre nuevo en `eco_nuevo`")
        print("   del activo que ya existe, y así conserva su historia.")
        print(f"   en el maestro y sin activo : {', '.join(e for _, e, _ in altas)}")
        print(f"   activo y sin código nuevo  : {', '.join(sorted(sueltos_base))}")

    # ── 4. CONFLICTOS de holograma ─────────────────────────────────────────────
    por_cod, por_eco = {}, {}
    for eco, cod in holo:
        if cod:
            por_cod.setdefault(norm_etiqueta(cod), set()).add(norm_eco(eco))
            por_eco.setdefault(norm_eco(eco), set()).add(norm_etiqueta(cod))
    choques = {c: e for c, e in por_cod.items() if len(e) > 1}
    dobles = {e: c for e, c in por_eco.items() if len(c) > 1}
    titulo(f"4. CONFLICTOS DE HOLOGRAMA ({len(choques) + len(dobles)}) · NO se vinculan")
    for c, ecos in choques.items():
        print(f"   un código en varios activos: {c} -> {', '.join(sorted(ecos))}")
    for e, cods in dobles.items():
        print(f"   un activo con varios códigos: {e} -> {', '.join(sorted(cods))}")
    if choques or dobles:
        print("   El holograma sustituye al económico tecleado; compartido, el escaneo es")
        print("   ambiguo y no se sabe a quién cargarle el diésel. Lo resuelve una persona.")

    # ── aplicar ────────────────────────────────────────────────────────────────
    if not aplicar:
        titulo("SIMULACIÓN · no se escribió nada")
        print("   Para aplicar altas y bajas:  --aplicar")
        db.close()
        return

    hechas = falladas = 0
    for ent, eco, placa in altas:
        ok, msg = catalogo.aplicar_propuesta(db, Propuesta("crear", ent, eco, placa or None))
        hechas += ok
        falladas += not ok
        if not ok:
            print(f"   ALTA FALLIDA {eco}: {msg}")

    revividas = 0
    for ent, obj, _ in revivir:
        ok, _m = catalogo.aplicar_propuesta(
            db, Propuesta("reactivar", ent, None, entidad_id=obj.id))
        revividas += ok

    ahora = datetime.now(timezone.utc)
    for _, obj, _, _, _, _ in bajas:
        obj.activo = False
        obj.verificado_en = ahora
    db.commit()

    titulo("APLICADO")
    print(f"   altas .......... {hechas}" + (f"  ({falladas} fallidas)" if falladas else ""))
    print(f"   reactivaciones . {revividas}")
    print(f"   bajas .......... {len(bajas)}")

    titulo("PARA REVERTIR LAS BAJAS")
    cu = [e for t, _, e, _, _, _ in bajas if t == "unidad"]
    cr = [e for t, _, e, _, _, _ in bajas if t == "remolque"]
    if cu:
        print("   UPDATE unidades SET activo=true WHERE clave IN ("
              + ", ".join(f"'{c}'" for c in cu) + ");")
    if cr:
        print("   UPDATE remolques SET activo=true WHERE eco IN ("
              + ", ".join(f"'{c}'" for c in cr) + ");")
    print("   Las altas se revierten a mano: son filas nuevas, y borrarlas sólo es seguro")
    print("   mientras nada las referencie.")
    db.close()


if __name__ == "__main__":
    a = sys.argv[1:]
    rutas = [x for x in a if not x.startswith("--")]
    main("--aplicar" in a,
         rutas[0] if len(rutas) > 0 else DEF_HOLO,
         rutas[1] if len(rutas) > 1 else DEF_PLAC)
