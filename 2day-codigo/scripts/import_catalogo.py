"""Importa el CATÁLOGO MAESTRO desde la hoja 'Operadores Motriz Remolque'.

Esa hoja es la fuente de verdad del cliente para tres cosas a la vez:
  · el padrón de operadores (con su NÚMERO consecutivo, que arranca en 1000),
  · el parque vehicular  (T### = tracto, C### = camión, numérico = remolque con thermo),
  · y la ASIGNACIÓN operador -> unidad.

Sirve además para acabar con los duplicados del padrón: el catálogo trae el nombre
canónico y completo ("APELLIDOS NOMBRE"), mientras que en la base fueron entrando
variantes sueltas desde WhatsApp ("LEOPOLDO SANCHEZ" vs "SANCHEZ ANDUJO LEOPOLDO").
El casado usa COINCIDENCIA TOTAL de palabras (un conjunto contenido en el otro), la misma
regla que el bot, así que da igual el orden y da igual si reportan por nombre o apellidos.

Por defecto hace SIMULACIÓN y no toca la base. Para escribir:
    python -m scripts.import_catalogo --aplicar
    python -m scripts.import_catalogo --aplicar --fusionar   (además fusiona duplicados)
"""

from __future__ import annotations

import argparse
import re
import sys
import unicodedata
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from sqlalchemy import select, update  # noqa: E402

from app.db import SessionLocal  # noqa: E402
from app.models import (  # noqa: E402
    AuditoriaThermo, Operador, Remolque, TipoUnidad, Unidad, Viaje,
)

HOJA = "Operadores Motriz Remolque"
NUMERO_BASE = 1000        # el primer operador del catálogo es el 1000; de ahí consecutivo

# Rendimiento objetivo por configuración (columna TIPO del catálogo), en km/L.
OBJETIVO = {"FULL": 1.9, "SENCILLO": 2.8, "THORTON": 4.0}

_STOP = {"DE", "DEL", "LA", "LAS", "LOS", "MC", "SAN"}


def tokens(nombre: str | None) -> set[str]:
    """Palabras significativas, sin acentos ni puntuación (igual que el bot)."""
    s = unicodedata.normalize("NFD", (nombre or "").upper())
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    s = re.sub(r"[^A-ZÑ0-9 ]", " ", s)
    return {t for t in s.split() if len(t) >= 3 and t not in _STOP}


def limpio(nombre: str | None) -> str:
    return " ".join(str(nombre or "").strip().upper().split())


def leer_catalogo(path: Path) -> list[dict]:
    import openpyxl

    ws = openpyxl.load_workbook(path, data_only=True)[HOJA]
    filas = []
    for r in range(2, ws.max_row + 1):
        op, uni, tipo = ws.cell(r, 1).value, ws.cell(r, 3).value, ws.cell(r, 4).value
        if not op and not uni:
            continue
        clave = "".join(str(uni).strip().upper().split()) if uni is not None else None
        filas.append({
            "fila": r,
            "operador": limpio(op) or None,
            # El número viene como fórmula encadenada (=B2+1) que openpyxl no evalúa:
            # se reconstruye por posición, que es exactamente lo que la fórmula hace.
            "numero": NUMERO_BASE + r - 2,
            "clave": clave,
            "tipo": limpio(tipo) or None,
        })
    return filas


def clase_unidad(clave: str | None) -> str | None:
    """T### -> tracto, C### -> camión, todo-dígitos -> remolque (caja con thermo)."""
    if not clave:
        return None
    if re.fullmatch(r"T\d+", clave):
        return "tracto"
    if re.fullmatch(r"C\d+", clave):
        return "camion"
    if clave.isdigit():
        return "remolque"
    return None


def importar(path: Path, aplicar: bool, fusionar: bool) -> dict:
    filas = leer_catalogo(path)
    s = SessionLocal()
    res = {
        "filas": len(filas), "unidades_nuevas": [], "remolques_nuevos": [],
        "operadores_nuevos": [], "operadores_renombrados": [], "numeros_asignados": 0,
        "asignaciones": 0, "objetivos": 0, "fusiones": [], "ambiguos": [], "sin_match": [],
        "numeros_liberados": [],
    }

    def numero_libre(num: int, para: Operador | None) -> bool:
        """Deja el número disponible para `para`. El catálogo del cliente manda: si lo
        ocupa alguien más, se le retira (y se reporta). Solo se hace cuando ese registro
        no tiene viajes — con historial de por medio se respeta y se reporta el choque."""
        dueno = s.execute(select(Operador).where(Operador.numero == num)).scalars().first()
        if dueno is None or (para is not None and dueno.id == para.id):
            return True
        con_viajes = s.execute(
            select(Viaje.id).where(Viaje.operador_id == dueno.id).limit(1)).first()
        if con_viajes:
            res["numeros_liberados"].append(
                {"numero": num, "conservado_por": dueno.nombre, "motivo": "tiene viajes"})
            return False
        dueno.numero = None
        s.flush()
        res["numeros_liberados"].append(
            {"numero": num, "retirado_a": dueno.nombre, "motivo": "sin viajes"})
        return True
    try:
        # ── Unidades y remolques ────────────────────────────────────────────
        for f in filas:
            cl = f["clave"]
            clase = clase_unidad(cl)
            if clase in ("tracto", "camion"):
                u = s.execute(select(Unidad).where(Unidad.clave == cl)).scalar_one_or_none()
                if u is None:
                    u = Unidad(clave=cl,
                               tipo=TipoUnidad.TRACTO if clase == "tracto" else TipoUnidad.CAMION)
                    s.add(u)
                    s.flush()
                    res["unidades_nuevas"].append(cl)
                obj = OBJETIVO.get(f["tipo"] or "")
                if obj and u.rendimiento_objetivo is None:
                    u.rendimiento_objetivo = obj
                    res["objetivos"] += 1
            elif clase == "remolque":
                r = s.execute(select(Remolque).where(Remolque.eco == cl)).scalar_one_or_none()
                if r is None:
                    s.add(Remolque(eco=cl, descripcion=f["tipo"]))
                    s.flush()
                    res["remolques_nuevos"].append(cl)

        # ── Operadores: el catálogo manda el nombre canónico y el número ────
        existentes = list(s.execute(select(Operador)).scalars())
        for f in filas:
            nom = f["operador"]
            if not nom:
                continue
            t = tokens(nom)
            exacto = next((o for o in existentes if limpio(o.nombre) == nom), None)
            cands = []
            if exacto is None and len(t) >= 2:
                for o in existentes:
                    ot = tokens(o.nombre)
                    if ot and min(len(t), len(ot)) >= 2 and (t <= ot or ot <= t):
                        cands.append(o)
            op = exacto
            if op is None and cands:
                nombres = {limpio(o.nombre) for o in cands}
                if len(nombres) > 1:
                    # Varias personas distintas casan: NO se toca, se reporta para revisión.
                    res["ambiguos"].append({"catalogo": nom, "candidatos": sorted(nombres)})
                    continue
                op = cands[0]
                if limpio(op.nombre) != nom:
                    res["operadores_renombrados"].append({"antes": op.nombre, "ahora": nom})
                    op.nombre = nom
            if op is None:
                num = f["numero"] if numero_libre(f["numero"], None) else None
                op = Operador(nombre=nom, numero=num, provisional=False)
                s.add(op)
                s.flush()
                existentes.append(op)
                res["operadores_nuevos"].append(nom)
            else:
                if op.numero != f["numero"] and numero_libre(f["numero"], op):
                    op.numero = f["numero"]
                    res["numeros_asignados"] += 1
                if op.provisional:
                    op.provisional = False       # está en el catálogo: ya es oficial
            f["op_id"] = op.id

            # Duplicados restantes: otros registros que son la MISMA persona
            if fusionar:
                for o in list(existentes):
                    if o.id == op.id:
                        continue
                    ot = tokens(o.nombre)
                    if ot and min(len(t), len(ot)) >= 2 and (t <= ot or ot <= t):
                        n = s.execute(select(Viaje.id).where(Viaje.operador_id == o.id)).all()
                        # Hay que repuntar TODAS las referencias antes de borrar: el operador
                        # cuelga de viajes, de la auditoría de thermo y de la unidad asignada.
                        s.execute(update(Viaje).where(Viaje.operador_id == o.id)
                                  .values(operador_id=op.id))
                        s.execute(update(AuditoriaThermo)
                                  .where(AuditoriaThermo.operador_id == o.id)
                                  .values(operador_id=op.id))
                        s.execute(update(Unidad).where(Unidad.operador_asignado_id == o.id)
                                  .values(operador_asignado_id=op.id))
                        res["fusiones"].append({"absorbido": o.nombre, "en": op.nombre,
                                                "viajes": len(n)})
                        s.flush()
                        s.delete(o)
                        s.flush()
                        existentes.remove(o)

        s.flush()

        # ── Asignación operador -> unidad ───────────────────────────────────
        for f in filas:
            if not f.get("op_id") or clase_unidad(f["clave"]) not in ("tracto", "camion"):
                continue
            u = s.execute(select(Unidad).where(Unidad.clave == f["clave"])).scalar_one_or_none()
            if u is not None and u.operador_asignado_id != f["op_id"]:
                u.operador_asignado_id = f["op_id"]
                res["asignaciones"] += 1

        # Operadores de la base que el catálogo no reconoce (altas del chat, bajas, etc.)
        cat_tok = [tokens(f["operador"]) for f in filas if f["operador"]]
        for o in s.execute(select(Operador)).scalars():
            ot = tokens(o.nombre)
            if not any(ct and ot and min(len(ct), len(ot)) >= 2 and (ct <= ot or ot <= ct)
                       for ct in cat_tok):
                res["sin_match"].append(o.nombre)

        if aplicar:
            s.commit()
        else:
            s.rollback()
    finally:
        s.close()
    return res


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("excel", nargs="?",
                    default=str(Path.home() / "Downloads" / "Bitacora 2025 2026.xlsx"))
    ap.add_argument("--aplicar", action="store_true", help="escribe en la base (si no, simula)")
    ap.add_argument("--fusionar", action="store_true",
                    help="además fusiona los duplicados que casan con el catálogo")
    a = ap.parse_args()

    r = importar(Path(a.excel), a.aplicar, a.fusionar)
    modo = "APLICADO" if a.aplicar else "SIMULACIÓN (nada se escribió)"
    print(f"=== Catálogo maestro · {modo} ===")
    print(f"filas leídas           : {r['filas']}")
    print(f"unidades nuevas        : {len(r['unidades_nuevas'])} {r['unidades_nuevas'][:8]}")
    print(f"remolques nuevos       : {len(r['remolques_nuevos'])} {r['remolques_nuevos'][:8]}")
    print(f"operadores nuevos      : {len(r['operadores_nuevos'])} {r['operadores_nuevos'][:5]}")
    print(f"nombres normalizados   : {len(r['operadores_renombrados'])}")
    for x in r["operadores_renombrados"][:8]:
        print(f"    {x['antes']!r} -> {x['ahora']!r}")
    print(f"números asignados      : {r['numeros_asignados']}")
    print(f"asignaciones a unidad  : {r['asignaciones']}")
    print(f"objetivos de rendim.   : {r['objetivos']}")
    print(f"fusiones de duplicados : {len(r['fusiones'])}")
    for x in r["fusiones"][:10]:
        print(f"    {x['absorbido']!r} -> {x['en']!r} ({x['viajes']} viajes)")
    print(f"números en conflicto   : {len(r['numeros_liberados'])}")
    for x in r["numeros_liberados"][:10]:
        print(f"    #{x['numero']}: {x}")
    print(f"AMBIGUOS (sin tocar)   : {len(r['ambiguos'])}")
    for x in r["ambiguos"][:10]:
        print(f"    catálogo {x['catalogo']!r} casa con {x['candidatos']}")
    print(f"en la base y NO en el catálogo: {len(r['sin_match'])}")
    print(f"    {r['sin_match'][:12]}")


if __name__ == "__main__":
    main()
