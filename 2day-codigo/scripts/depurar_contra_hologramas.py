"""Deja el catálogo alineado con el maestro de hologramas. Decisión del dueño (3-sep-2026):
"fuera de estas no existen más unidades".

QUÉ HACE, y qué NO:
  · Activos vigentes que el maestro no lista  -> `activo = False`. NO se borran: conservan sus
    viajes, sus asientos y sus etiquetas, y volver atrás es un UPDATE.
  · Etiquetas cuyo código ya no está en el maestro -> `estado = 'retirada'`. Tampoco se borran:
    las cargas que se registraron mientras estuvieron vigentes necesitan a qué apuntar.

LA REGLA QUE PROTEGE, y por qué existe: **un activo con consumo NO se da de baja**, aunque el
maestro no lo liste. Ya pasó con 5311801 y 5311808, que queman diésel sin aparecer en PLACAS.
Aquí vuelve a pasar con T203: tiene litros atribuidos en el libro mayor, y darlo de baja los
dejaría sin dueño. El script lo detecta solo y lo deja fuera, diciéndolo.

Uso:
    python -m scripts.depurar_contra_hologramas --dry-run   (no escribe NADA)
    python -m scripts.depurar_contra_hologramas --aplicar
"""

import io as _io
import os
import sys

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from datetime import datetime, timezone

from sqlalchemy import select

from app.catalogo import norm_eco, norm_etiqueta
from app.db import SessionLocal
from app.models import AsientoConsumo, EtiquetaActivo, Remolque, Unidad, Viaje

DEF = os.path.join(os.path.expanduser("~/Downloads"), "HOLOGRAMA (1).xlsx")


def leer(path):
    with open(path, "rb") as fh:
        wb = openpyxl.load_workbook(_io.BytesIO(fh.read()), data_only=True, read_only=True)
    hoja = wb[wb.sheetnames[0]]
    ecos, cods = set(), set()
    for r in hoja.iter_rows(min_row=4, values_only=True):
        if len(r) > 1 and r[1]:
            ecos.add(norm_eco(r[1]))
        if len(r) > 2 and r[2]:
            cods.add(norm_etiqueta(r[2]))
    wb.close()
    return ecos, cods


def consumo(db, obj, es_unidad):
    q = db.query(AsientoConsumo).filter(
        AsientoConsumo.unidad_id == obj.id if es_unidad
        else AsientoConsumo.remolque_id == obj.id)
    filas = q.all()
    return len(filas), sum(float(a.litros or 0) for a in filas)


def main(aplicar: bool, path: str):
    if not os.path.exists(path):
        print(f"NO SE ENCUENTRA: {path}")
        sys.exit(1)
    ecos, cods = leer(path)
    db = SessionLocal()

    print("=" * 78)
    print(f"MAESTRO: {os.path.basename(path)} · {len(ecos)} activos · {len(cods)} códigos")
    print("-" * 78)

    bajas, protegidos = [], []
    for u in db.execute(select(Unidad).where(Unidad.activo.is_(True))).scalars():
        if norm_eco(u.clave) in ecos:
            continue
        n, lts = consumo(db, u, True)
        nv = db.query(Viaje).filter(Viaje.unidad_id == u.id).count()
        (protegidos if n else bajas).append(("unidad", u, nv, n, lts))
    for r in db.execute(select(Remolque).where(Remolque.activo.is_(True))).scalars():
        if ({norm_eco(r.eco), norm_eco(r.eco_nuevo)} - {""}) & ecos:
            continue
        n, lts = consumo(db, r, False)
        nv = db.query(Viaje).filter(Viaje.remolque_thermo == r.eco).count()
        (protegidos if n else bajas).append(("remolque", r, nv, n, lts))

    sobran = [e for e in db.execute(select(EtiquetaActivo)).scalars()
              if e.codigo not in cods and e.estado != "retirada"]

    print(f"\n  SE DAN DE BAJA ({len(bajas)}) · el maestro no los lista y no tienen consumo")
    for tipo, o, nv, _, _ in bajas:
        print(f"     {(o.clave if tipo == 'unidad' else o.eco):<12} {tipo:<9} "
              f"{nv:>4} viajes (se conservan)")

    if protegidos:
        print(f"\n  NO SE TOCAN ({len(protegidos)}) · tienen litros atribuidos en el libro mayor")
        for tipo, o, nv, n, lts in protegidos:
            print(f"     {(o.clave if tipo == 'unidad' else o.eco):<12} {tipo:<9} "
                  f"{n:>4} asientos · {lts:>9,.1f} L   <- darlo de baja dejaría estos litros sin dueño")

    print(f"\n  ETIQUETAS QUE SE RETIRAN ({len(sobran)}) · su código ya no está en el maestro")
    for e in sobran:
        print(f"     {e.codigo:<16} {e.eco_texto or '-':<10} (estaba '{e.estado}')")

    if not aplicar:
        print("\n" + "=" * 78)
        print("SIMULACIÓN: no se escribió nada. Para aplicar: --aplicar")
        db.close()
        return

    ahora = datetime.now(timezone.utc)
    for _, o, _, _, _ in bajas:
        o.activo = False
    for e in sobran:
        e.estado = "retirada"
        e.nota = "retirada: su código ya no aparece en el maestro de hologramas"
        e.actualizada_en = ahora
    db.commit()

    print("\n" + "=" * 78)
    print(f"{len(bajas)} activos dados de baja · {len(sobran)} etiquetas retiradas.")
    print("Ninguna fila se borró.")
    u = [o.clave for t, o, _, _, _ in bajas if t == "unidad"]
    r = [o.eco for t, o, _, _, _ in bajas if t != "unidad"]
    print("\nRevertir:")
    if u:
        print("  UPDATE unidades  SET activo=true WHERE clave IN ("
              + ", ".join(f"'{x}'" for x in u) + ");")
    if r:
        print("  UPDATE remolques SET activo=true WHERE eco IN ("
              + ", ".join(f"'{x}'" for x in r) + ");")
    if sobran:
        print("  UPDATE etiquetas_activo SET estado='vinculada' WHERE codigo IN ("
              + ", ".join(f"'{e.codigo}'" for e in sobran) + ");   -- revisar estado previo")
    db.close()


if __name__ == "__main__":
    a = sys.argv
    ruta = next((x for x in a[1:] if not x.startswith("--")), DEF)
    main("--aplicar" in a, ruta)
